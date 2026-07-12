"""
capital_gains.py -- classify equity sale lots into STCG (Sec 111A) vs LTCG (Sec 112A), apply
31-Jan-2018 grandfathering, and produce filing-ready splits + a summary.

WHY THIS EXISTS
  At return time you reconcile three views of the same equity sales:
    * AIS (tax dept)   -- lists every sale + consideration, but the COST is often ZERO for old / off-market
                          lots (the dept doesn't know your purchase price) -> you must fill it in.
    * Broker P&L       -- has the real buy value + buy/sell dates -> the source that fills those zero costs.
    * Your book (Tally)-- the equity ledger/stock the sales came from.
  Once you have a clean, normalised lot list (buy value + dates filled in), the mechanical part is:
  split by holding period, grandfather the pre-2018 long-term lots, and total up for Schedule 111A/112A.
  This script does that mechanical part. Building the normalised sheet (reconciling AIS <-> broker) is the
  human/AI step the sample lesson teaches.

INPUT  (CSV or XLSX; one row per sale lot)  -- headers are case/space-insensitive:
    security, isin, buy_date, buy_value, sell_date, sell_value, fmv_31jan2018 (optional)
  * values are TOTALS in rupees (buy_value = total cost of that lot; sell_value = total sale consideration;
    fmv_31jan2018 = qty x highest quoted price on 31-Jan-2018, only needed for grandfathered long-term lots).
  * dates accept YYYY-MM-DD, DD-MM-YYYY, or DD-MMM-YYYY.

USAGE
    python capital_gains.py --file trades.csv                       # print the split + summary
    python capital_gains.py --file trades.csv --out gains.xlsx      # also write a filing-ready workbook

RULES APPLIED (listed equity / equity MF, STT paid)  -- verify against the Finance Act for your year:
  * Long-term if held > 12 months (default 365 days; --long-term-days to change); else short-term.
  * STCG 111A = sale - actual cost.
  * LTCG 112A = sale - deemed cost. For lots BOUGHT BEFORE 01-Feb-2018, grandfathering applies:
        deemed cost = HIGHER of ( actual cost , LOWER of ( FMV on 31-Jan-2018 , sale consideration ) )
    (so the FMV can only reduce a gain, never manufacture a loss). Needs the fmv_31jan2018 column.
  * Summary shows total LTCG less the Sec 112A exemption (default Rs 1,25,000; --ltcg-exemption to change).

NOT TAX ADVICE. Rates, the exemption, and the holding-period rule are set by law and change year to year;
this classifies and totals -- your CA signs off. Grandfathering only affects long-term lots bought pre-2018.
"""
import sys, csv, argparse, datetime as dt

FIELDS = ["security", "isin", "buy_date", "buy_value", "sell_date", "sell_value", "fmv_31jan2018"]


def _norm(h):
    return (h or "").strip().lower().replace(" ", "_").replace("-", "_")


def parse_date(s):
    s = (s or "").strip()
    if not s:
        raise ValueError("empty date")
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d-%b-%Y", "%d-%B-%Y", "%Y/%m/%d"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"unrecognised date: {s!r} (use YYYY-MM-DD / DD-MM-YYYY / DD-MMM-YYYY)")


def num(s):
    s = str(s or "").strip().replace(",", "").replace("Rs.", "").replace("₹", "")
    return float(s) if s else 0.0


def read_rows(path):
    if path.lower().endswith((".xlsx", ".xlsm")):
        try:
            import openpyxl
        except ImportError:
            raise SystemExit("openpyxl not installed for .xlsx input -- `pip install openpyxl`, or pass a CSV.")
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        hdr = [_norm(str(c)) for c in rows[0]]
        return [dict(zip(hdr, r)) for r in rows[1:] if any(c is not None for c in r)]
    with open(path, newline="", encoding="utf-8-sig") as f:
        rdr = csv.DictReader(f)
        rdr.fieldnames = [_norm(h) for h in (rdr.fieldnames or [])]
        return [row for row in rdr if any((v or "").strip() for v in row.values())]


def classify(rows, long_term_days, grandfather_before):
    out = []
    for i, r in enumerate(rows, 1):
        r = {_norm(k): v for k, v in r.items()}
        sec = (r.get("security") or "").strip()
        isin = (r.get("isin") or "").strip()
        buy_d = parse_date(r.get("buy_date"))
        sell_d = parse_date(r.get("sell_date"))
        buy_v = num(r.get("buy_value"))
        sell_v = num(r.get("sell_value"))
        fmv = num(r.get("fmv_31jan2018"))
        if sell_d < buy_d:
            raise SystemExit(f"row {i} ({sec}): sell date {sell_d} is before buy date {buy_d}")
        held = (sell_d - buy_d).days
        is_long = held > long_term_days
        grandfathered = False
        if is_long and buy_d < grandfather_before and fmv > 0:
            deemed_cost = max(buy_v, min(fmv, sell_v))
            grandfathered = True
        else:
            deemed_cost = buy_v
        gain = round(sell_v - deemed_cost)
        out.append({
            "#": i, "security": sec, "isin": isin,
            "buy_date": buy_d.isoformat(), "sell_date": sell_d.isoformat(),
            "held_days": held, "term": "LTCG" if is_long else "STCG",
            "itr_section": "112A" if is_long else "111A",
            "buy_value": round(buy_v), "sale_consideration": round(sell_v),
            "deemed_cost": round(deemed_cost), "gain_loss": gain,
            "grandfathered": "yes" if grandfathered else "",
        })
    return out


def summarise(rows, ltcg_exemption):
    st = [r for r in rows if r["term"] == "STCG"]
    lt = [r for r in rows if r["term"] == "LTCG"]
    st_gain = sum(r["gain_loss"] for r in st)
    lt_gain = sum(r["gain_loss"] for r in lt)
    lt_taxable = max(0, round(lt_gain - ltcg_exemption))
    return {
        "stcg_lots": len(st), "stcg_111a_net": round(st_gain),
        "ltcg_lots": len(lt), "ltcg_112a_net": round(lt_gain),
        "ltcg_exemption": ltcg_exemption, "ltcg_112a_taxable_after_exemption": lt_taxable,
        "grandfathered_lots": sum(1 for r in lt if r["grandfathered"]),
    }


def write_xlsx(path, rows, summ):
    try:
        import openpyxl
    except ImportError:
        raise SystemExit("openpyxl not installed -- `pip install openpyxl`, or drop --out to just print.")
    wb = openpyxl.Workbook()
    cols = ["#", "security", "isin", "buy_date", "sell_date", "held_days", "term", "itr_section",
            "buy_value", "sale_consideration", "deemed_cost", "gain_loss", "grandfathered"]

    def sheet(title, subset):
        ws = wb.create_sheet(title)
        ws.append(cols)
        for r in subset:
            ws.append([r[c] for c in cols])

    wb.remove(wb.active)
    sheet("STCG_111A", [r for r in rows if r["term"] == "STCG"])
    sheet("LTCG_112A", [r for r in rows if r["term"] == "LTCG"])
    sheet("All_Lots", rows)
    ws = wb.create_sheet("Summary", 0)
    ws.append(["Capital gains split (Sec 111A / 112A) -- verify rates & thresholds vs the Finance Act"])
    ws.append([])
    for k, v in summ.items():
        ws.append([k.replace("_", " "), v])
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description="Split equity sales into STCG 111A / LTCG 112A (+ grandfathering).")
    ap.add_argument("--file", required=True, help="normalised lot list (CSV/XLSX)")
    ap.add_argument("--out", help="write a filing-ready workbook (.xlsx)")
    ap.add_argument("--long-term-days", type=int, default=365, help="holding-period cutoff (default 365 = >12 months)")
    ap.add_argument("--grandfather-before", default="2018-02-01", help="grandfather lots bought before this date")
    ap.add_argument("--ltcg-exemption", type=int, default=125000, help="Sec 112A exemption for the summary")
    a = ap.parse_args()

    rows = classify(read_rows(a.file), a.long_term_days, parse_date(a.grandfather_before))
    if not rows:
        raise SystemExit("no rows found -- check the headers: " + ", ".join(FIELDS))
    summ = summarise(rows, a.ltcg_exemption)

    w = max(len(r["security"]) for r in rows) + 1
    print(f"\n{'security':<{w}} {'term':<5} {'sec':<5} {'sale':>12} {'cost':>12} {'gain/loss':>12}  note")
    print("-" * (w + 55))
    for r in sorted(rows, key=lambda x: (x["term"] != "STCG", x["security"])):
        note = "grandfathered" if r["grandfathered"] else ""
        print(f"{r['security']:<{w}} {r['term']:<5} {r['itr_section']:<5} "
              f"{r['sale_consideration']:>12,} {r['deemed_cost']:>12,} {r['gain_loss']:>12,}  {note}")
    print("\nSUMMARY")
    print(f"  STCG (111A): {summ['stcg_lots']} lots, net {summ['stcg_111a_net']:>12,}")
    print(f"  LTCG (112A): {summ['ltcg_lots']} lots, net {summ['ltcg_112a_net']:>12,}"
          f"  ({summ['grandfathered_lots']} grandfathered)")
    print(f"    less Sec 112A exemption {summ['ltcg_exemption']:,} -> taxable LTCG "
          f"{summ['ltcg_112a_taxable_after_exemption']:>12,}")
    if a.out:
        write_xlsx(a.out, rows, summ)
        print(f"\nwrote {a.out}. VERIFY totals against your broker P&L and AIS before filing.")


if __name__ == "__main__":
    main()
