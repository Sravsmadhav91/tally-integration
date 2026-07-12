"""
build_lesson0.py -- generate the Lesson 0 (basics) source files:
  * lesson0_transactions.xlsx  -- 4 transactions to post as vouchers (Receipt/Payment/Contra/Journal)
  * lesson0_bank_page.pdf      -- a 3-line mini bank statement to post as bank entries

These are the "learn to enter data" warm-up BEFORE the reconciliation exercises. They are dated
Jul-Sep 2025 (outside the Apr-Jun reconciliation window) so they don't disturb the sandbox, and
every date is the 1st/2nd/last of a month (works in Educational mode). Re-run to regenerate.

Needs: openpyxl (core) and pymupdf (optional power tool -- pip install pymupdf).
"""
import os
HERE = os.path.dirname(os.path.abspath(__file__))

# (Date, Voucher Type, Particulars, Amount, Debit ledger, Credit ledger)
TXNS = [
    ("02-07-2025", "Payment", "Office rent - July",                 25000, "Office Rent",     "Practice Bank"),
    ("31-07-2025", "Receipt", "Consulting fee - Client ABC",        55000, "Practice Bank",   "Consulting Income"),
    ("01-08-2025", "Contra",  "Transfer to Practice Bank 2",        10000, "Practice Bank 2", "Practice Bank"),
    ("31-08-2025", "Journal", "Electricity bill received (unpaid)",  3000, "Electricity",     "Vendor XYZ"),
]

# Mini bank statement (Practice Bank). (Date, Description, Debit, Credit)
BANK = [
    ("01-09-2025", "NEFT-CLIENT PQR",   0,    40000),
    ("02-09-2025", "ELECTRICITY BOARD", 3000, 0),
    ("30-09-2025", "BANK CHARGES",      150,  0),
]


def build_xlsx():
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Transactions to enter"
    hdr = ["Date", "Voucher Type", "Particulars", "Amount", "Debit (Dr)", "Credit (Cr)"]
    for i, h in enumerate(hdr, 1):
        c = ws.cell(1, i, h); c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for r, t in enumerate(TXNS, 2):
        for i, val in enumerate(t, 1):
            ws.cell(r, i, val)
        ws.cell(r, 4).number_format = "#,##0"
    widths = [12, 14, 34, 10, 16, 18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.cell(len(TXNS) + 3, 1, "Post each row as the named voucher type. The Dr/Cr columns tell you the two ledgers.")
    wb.save(os.path.join(HERE, "lesson0_transactions.xlsx"))
    print("Wrote lesson0_transactions.xlsx")


def build_pdf():
    try:
        import fitz  # pymupdf
    except ImportError:
        print("SKIP lesson0_bank_page.pdf -- install pymupdf (pip install pymupdf) then re-run.")
        return
    doc = fitz.open(); page = doc.new_page(width=595, height=842)
    y = [60]

    def line(txt, size=11, dy=18, bold=False):
        page.insert_text((60, y[0]), txt, fontsize=size, fontname="hebo" if bold else "helv")
        y[0] += dy

    line("PRACTICE BANK", 16, 24, bold=True)
    line("Account statement (learning sample) -- Account: Practice Bank XXXX1234", 10, 22)
    line("Date          Description                          Debit        Credit", 11, 6, bold=True)
    line("-" * 74, 11, 18)
    for d, desc, dr, cr in BANK:
        drs = f"{dr:,}" if dr else ""
        crs = f"{cr:,}" if cr else ""
        line(f"{d}    {desc:<34} {drs:>10} {crs:>10}", 11, 18)
    line("-" * 74, 11, 22)
    line("Post each line as a bank voucher: money IN = Receipt, money OUT = Payment.", 10, 16)
    line("Match the other ledger from the description (client -> income, electricity -> expense, etc.).", 10, 16)
    doc.save(os.path.join(HERE, "lesson0_bank_page.pdf")); doc.close()
    print("Wrote lesson0_bank_page.pdf")


if __name__ == "__main__":
    build_xlsx()
    build_pdf()
